# -*- coding: utf-8 -*-
r"""生成《义符——基于母语义类系统的外语学习软件规划书》，照公文格式。
规格沿用 D:\SRT项目\_reformat_gongwen.py（用户 2026-10-05 定死的那套）。"""
import os, io, sys, re, zipfile
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

OUT_DIR = r'D:\AI coding\规划书'   # 2026-10-06 用户定：规划书统一放汇总目录
DST = os.path.join(OUT_DIR, '义符_产品规划书_v1.docx')

FS, HT, KT, HZ = '仿宋', '黑体', '楷体', '华文中宋'
SZ = 16          # 三号
LS = 28          # 行距固定值 28 磅
TBLZ = 14        # 表格四号
TBL_LS = 22
WEST = 'Times New Roman'
TEXTW = Cm(15.6)

# ============================================================
# 内容定义
#   ('title'|'subtitle'|'info'|'h0'|'h1'|'h2'|'p'|'li'|'caption'|'table'|'pagebreak', payload)
# ============================================================
C = []
A = C.append

# ---------- 封面 ----------
A(('title', '义符'))
A(('subtitle', '——基于母语义类系统的多端外语学习软件规划书'))
A(('info', '文档版本：v1.0'))
A(('info', '编制日期：2026年10月6日'))
A(('info', '文档性质：产品规划书（立项前论证）'))
A(('blank', ''))

# ---------- 一 ----------
A(('h0', '一、项目背景与问题定义'))

A(('h1', '（一）一个被浪费的红利：汉语母语者的义类直觉'))
A(('p', '中国学生学英语，普遍从初中开始、持续十年以上，投入的总时长以千小时计，但绝大多数人最终停在"能读不能说、能背不能用"的状态。归因通常指向词汇量和语法，但更深的问题在于：母语在这套学习过程中始终被当成"翻译中介"，而不是"认知脚手架"。'))
A(('p', '当一个中国学生看到 apple，他的学习路径是 apple→苹果→那个红色水果。这是一条两次跳跃的链路：先用外文符号唤起中文符号，再由中文符号唤起概念。母语在这里是负担——它增加了一跳，而不是减少一跳。'))
A(('p', '而汉语母语者本应拥有一个其他语种学习者没有的优势：汉语是当今世界唯一仍在日常使用的大规模表意文字系统。每一个受过基础教育的中国人，脑子里都装着一套运行了十几年、完全自动化的"义类识别器"。'))

A(('h1', '（二）义符：汉语母语者被忽视的认知资产'))
A(('p', '汉字由部件构成，部件分两类：表义的叫形旁（义符），表音的叫声符。看到"氵"，不需要认识整个字，就能判断它和水有关；看到"扌"，就知道和手部动作有关；看到"忄"，就知道和情绪有关。这个判断是瞬时的、无意识的、几乎不消耗记忆负荷的。'))
A(('p', '这就是义符：一个固定的字形部件，携带一个稳定的语义类属，可以复用到成百上千个字上。一个高中毕业生熟练掌握的义符大约在一百到两百个之间，覆盖常用字的大部分。这套系统是免费的、已经装好的、不需要重新学习的。'))
A(('p', '问题在于：它从没被用来对接外语。'))

A(('h1', '（三）三个断层'))
A(('li', '断层一：义符能力与外语学习完全脱节。学生在语文课上学的"氵表示水、扌表示手"，和英语课上背的单词，是两套永不交汇的知识。没人告诉他 aquarium 的 aqua 和"江"的三点水是同一个义类。'))
A(('li', '断层二：单词被当作不可分的整体记忆。"manufacture"在多数学生眼里是一串字母，靠重复抄写记下来。但它等于 manu（手，拉丁 manus）加 fact（做，拉丁 facere）加 ure（名词后缀），意思是"用手做出来的东西"。这个拆解一旦建立，不仅记住一个词，还顺带激活了 manual、manuscript、manicure、factory、fact、perfect、defect 一整族词。'))
A(('li', '断层三：典故与词源被当成趣味知识，而非记忆基础设施。为什么牛肉叫 beef，牛却叫 cow？因为 1066 年诺曼征服后，说法语的贵族吃牛肉、说英语的盎格鲁-撒克逊人养牛。活着的动物用英语词，端上餐桌的肉用法语词，于是有了 cow/beef、pig/pork、sheep/mutton、calf/veal 这一组对照。这不是八卦，这是一条能一次性锁死十个单词、并且不容易忘的记忆线。但课堂几乎从不讲。'))

A(('h1', '（四）问题定义'))
A(('p', '本项目要解决的问题可以精确表述为：如何把汉语母语者已有的义符系统，与外语的词根词缀系统建立显式映射，使母语从学习的负担转为学习的杠杆。'))
A(('p', '这个问题背后有一个清晰的、可检验的假设：'))
A(('p', '当学习者在学习一个外文词时，同时获得该词的"义类归属"和"汉语同义类的义符对照"，其记忆保持率与猜词迁移能力，是否显著高于仅提供中文释义的对照组。'))
A(('p', '这个假设是整个项目的靶心，也是需要最先被验证的东西。在它被验证之前，本规划书中的一切功能设计都只是待检验的猜想。'))

# ---------- 二 ----------
A(('h0', '二、核心洞察：两套义类系统的同构性'))

A(('h1', '（一）从词根记忆法到双系统对齐'))
A(('p', '以词根词缀记单词不是新事物。市面上的词根词典与各类考纲词汇书，逻辑都是"把单词拆成词根"，这是有效的，但存在一个根本缺陷：它是用英语解释英语。'))
A(('p', '学生记 manu 等于 hand，仍然是拿一个陌生的符号去对应另一个陌生的符号。他必须先把 hand 记住，才能用 hand 去记 manu。这是叠罗汉，地基没有变厚。'))
A(('p', '本项目的主张是：用母语的义符去锚定外文的词根。manu 不是对应 hand，而是对应"扌"。扌这个部件，学习者已经用了十几年，是全自动的。当 manu 被挂到"扌"上时，它挂到的不是一个新记忆，而是一个已经存在的、极其牢固的认知结构。'))

A(('h1', '（二）同构性论证'))
A(('p', '两套系统在功能层面上是同构的，对照如下：'))
A(('caption', '表1　汉语义符系统与英语词根系统的结构对照'))
A(('table', [
    ['维度', '汉语义符系统', '英语词根词缀系统'],
    ['载体', '字形部件（形旁）', '拼写片段（词根与前后缀）'],
    ['功能', '提示语义类属', '提示语义类属'],
    ['数量级', '常用义符约100至200个', '常用词根约400至600个，词缀约100个'],
    ['复用率', '单个义符覆盖数十至数百字', '单个词根覆盖数个至数十词'],
    ['来源', '本土演化（象形、隶变、简化）', '外来借入（拉丁、希腊为主，日耳曼为底）'],
    ['位置', '相对固定（左、右、上、下）', '相对自由'],
    ['学习时点', '小学阶段随母语习得完成', '从未被系统教授'],
]))
A(('p', '必须诚实说明的差异：两套系统并非一一对应。汉语义符是本土演化的产物，位置相对固定；英语词根大量来自拉丁与希腊借词，且一个义类可能同时有拉丁来源和希腊来源两个词根并存，例如表示"水"的拉丁 aqua- 与希腊 hydr-。因此映射不是"翻译"，而是"归类"——多个外文词根可以归入同一个义类，正如多个汉字可以共用一个义符。'))
A(('p', '这个差异不是缺陷，反而是产品的一个额外价值点：它让学生理解为什么英语里同一个意思会有两套词根。aquarium 和 hydrogen 都是水，因为一个是拉丁来的日常词，一个是希腊来的学术词。这本身就是一段值得讲的历史。'))

A(('h1', '（三）映射表样例'))
A(('p', '以下是"义类、汉语义符、英语词根"三层映射的核心样例，完整首批见附录A：'))
A(('caption', '表2　义类映射表样例（首批12条）'))
A(('table', [
    ['义类', '汉语义符', '英语词根（来源）', '例词'],
    ['水·液体', '氵 水', 'aqua-（拉）, hydr-（希）, mar-（拉）', 'aquarium, hydrogen, marine'],
    ['手·操作', '扌 手', 'manu-（拉）, chiro-（希）', 'manual, manufacture, chiropractor'],
    ['心·情感', '忄 心', 'cord-（拉）, psych-（希）', 'courage, cordial, psychology'],
    ['走·行进', '辶 彳 足', 'ped-（拉）, gress-（拉）, vad-（拉）', 'pedal, progress, invade'],
    ['言·言语', '讠 言', 'loqu-（拉）, dict-（拉）, log-（希）', 'eloquent, predict, dialogue'],
    ['生·生命', '生', 'bio-（希）, viv-（拉）', 'biology, survive, revive'],
    ['火·光热', '火 灬', 'ign-（拉）, photo-（希）', 'ignite, photograph'],
    ['土·大地', '土', 'terr-（拉）, geo-（希）', 'territory, geology'],
    ['石·坚硬', '石', 'lith-（希）, petr-（希）', 'monolith, petrify'],
    ['大·巨大', '大', 'magn-（拉）', 'magnify, magnitude'],
    ['小·微小', '小', 'micro-（希）', 'microscope, microbe'],
    ['目·观看', '目', 'vid-/vis-（拉）, spec-（拉）', 'visible, inspect, perspective'],
]))
A(('p', '需要说明一个常见陷阱，它也应当成为产品内容的一部分：ped- 有两个来源完全不同的同形词根。来自拉丁 pes（脚）的 ped- 见于 pedal、pedestrian；来自希腊 pais（儿童）的 ped- 见于 pediatric、pedagogy。两者拼写相同、来源不同、意思无关。这类"假同源"是学习者最容易踩的坑，也正是产品可以体现专业度的地方。'))

A(('h1', '（四）为什么这是母语逻辑，而不只是记忆技巧'))
A(('p', '记忆技巧是外加的，需要刻意维持。母语逻辑是调用已经自动化的认知结构。'))
A(('p', '汉语母语者读"江"这个字，不需要先想"这是水"再理解，三点水已经直接参与认知了。当他把这个自动化机制挪用到 aquarium 上，他用的不是一条需要记住的规则，而是一套已经内化的操作系统。'))
A(('p', '这是本项目与所有词根记忆法产品的本质区别：别人教规则，我们调用本能。'))

# ---------- 三 ----------
A(('h0', '三、产品定位与目标用户'))

A(('h1', '（一）一句话定位'))
A(('p', '义符是一个帮助汉语母语者用"认偏旁"的直觉来"认词根"的多端学习工具，让学外语从记符号变成认结构。'))

A(('h1', '（二）产品名'))
A(('p', '暂定名义符。义符即形旁，与声符相对，是汉字学的正式术语，准确指向产品的核心机制。备选方案如下：'))
A(('caption', '表3　产品名备选方案'))
A(('table', [
    ['备选名', '取义', '优劣'],
    ['义符', '直接对应核心机制', '精准，但稍显学术'],
    ['训诂', '中国传统语文学，专研字义源流', '文化感强，但对英语场景略绕'],
    ['同源', '两个系统本出同理', '简洁，但易与"同源词"概念混淆'],
    ['双根', '义符与词根', '好记，但易与数学"双根"混淆'],
    ['字根词根', '直白描述', '好懂，不够精炼'],
]))

A(('h1', '（三）目标用户'))
A(('caption', '表4　目标用户分层'))
A(('table', [
    ['用户群', '特征', '核心诉求'],
    ['备考型大学生', '四六级、考研英语，词汇量大但死记硬背', '提高背词效率，减少遗忘'],
    ['英语中级学习者', '能读不能猜，遇生词即卡', '获得猜词与推义能力'],
    ['语言兴趣者', '对词源、历史典故有兴趣', '内容深度与趣味'],
    ['语文教育关注者', '关心汉字文化传承', '反向加深对母语的理解'],
    ['对外汉语教学者', '教外国人学汉字', '同一引擎可反向使用，见附录D'],
]))

A(('h1', '（四）与现有产品的差异'))
A(('p', '市面同类产品可分为三类，均未覆盖本项目的核心机制：'))
A(('caption', '表5　竞品覆盖度对照'))
A(('table', [
    ['类别', '代表形态', '覆盖什么', '缺什么'],
    ['词根词缀背词', '各类词根记忆应用、词根词典', '英语词根拆解', '与母语无连接，仍是"用英语记英语"'],
    ['汉字字源工具', '字源查询应用、汉字演变演示', '汉字甲骨文到楷书的演变', '只服务母语，不涉外语'],
    ['通用背词应用', '主流单词记忆应用', '词库规模、复习算法', '记忆挂钩只有中文释义与图片'],
]))
A(('p', '空白点正是"对齐"这一层：把汉语义符与英语词根挂到同一个义类节点上，用母语已内化的机制去驱动外语学习。这一层目前没有成熟产品。'))

# ---------- 四 ----------
A(('h0', '四、核心数据架构：义类映射引擎'))

A(('h1', '（一）三层数据模型'))
A(('p', '整个产品的数据基础是三层结构。'))
A(('li', '第一层：义类。与语言无关的语义原子，例如水、手、心、走、言。这是全库的骨架，每个义类是一个节点。'))
A(('li', '第二层：义符与词根。义类在具体语言中的表现形式。汉语侧是义符，如氵、扌、忄；英语侧是词根词缀，如 aqua-、manu-、cord-。一个义类可以挂多个义符或词根。'))
A(('li', '第三层：词条。具体语言单位。汉语侧是字，如江、河、湖、海；英语侧是词，如 aquarium、marine、hydrate。每个词条挂到它的义符或词根上。'))
A(('p', '三层是树状到网状的混合结构：义类到标记到词条为树，词条之间因共享典故、共享来源语言而互连为网。'))

A(('h1', '（二）词条数据模型'))
A(('p', '每个外文词条包含以下字段，完整字段表见附录B。'))
A(('caption', '表6　词条数据模型示例（以 manufacture 为例）'))
A(('table', [
    ['字段', '说明', '示例'],
    ['拼写', '标准拼写', 'manufacture'],
    ['音标', '国际音标', '/ˌmænjuˈfæktʃə(r)/'],
    ['义类', '主义类标识', '义类_手'],
    ['词根分解', '部件序列及各自含义', 'manu（手）+ fact（做）+ ure（名词后缀）'],
    ['汉语对应', '同义类的汉字示例', '扌部：打、做、造、操'],
    ['典故', '词源故事', 'manu 源自拉丁 manus（手），与 manual、manuscript 同源'],
    ['来源语言', '拉丁、希腊或日耳曼', '拉丁'],
    ['例句', '双语例句', '略'],
    ['难度等级', '对应考纲等级', '四级'],
    ['关联词', '同根词列表', 'manual, manuscript, manicure, factory, fact'],
]))

A(('h1', '（三）典故数据库'))
A(('p', '典故是产品的深度来源，也是留存的关键。数据库分两条线。'))
A(('p', '汉语线记录字的字形演变史，即甲骨文、金文、小篆、隶变、楷书、简化这一条脉络。例如"水"在甲骨文中象流水之形，"手"象五指之形。'))
A(('p', '英语线记录词根进入英语的历史路径，主要有四条管道。'))
A(('li', '日耳曼底层：古英语原生词，构成日常基础词汇，如 cow、house、water。'))
A(('li', '维京侵入（9至11世纪）：古诺斯语贡献了大量基础词，如 sky、skin、egg、they、them。'))
A(('li', '诺曼征服（1066年）：诺曼法语涌入，形成活物用英语、餐桌用法语的分层，如 cow/beef、pig/pork、sheep/mutton。'))
A(('li', '文艺复兴（15至16世纪）：希腊与拉丁词根大规模进入学术词汇，如 biology、anthropology、philosophy。'))
A(('p', '两条线的并置本身就是产品最独特的内容：同一个历史分层现象，汉语和英语各自发生过一次。'))

A(('h1', '（四）多语种扩展接口'))
A(('p', '按架构预留多语种口子的要求，关键在于第一层的义类节点不绑定任何语言。若未来加入其他语种，可复用情况如下。'))
A(('caption', '表7　多语种扩展的复用关系'))
A(('table', [
    ['语种', '义符或词根系统', '与汉语的关系', '复用度'],
    ['日语', '汉字（音读与训读）加假名', '同源但同形异义，需额外处理"陷阱"', '中'],
    ['法语', '拉丁词根', '与英语大量共享拉丁部分', '高'],
    ['德语', '日耳曼词根加复合构词', '与英语同属日耳曼语族，可交叉对照', '中'],
    ['西班牙语', '拉丁词根', '与法语同理', '高'],
]))
A(('p', '因此数据模型从第一天起就不应把"英语"写死，而应以"语言、标记、词条"的多语言结构存储。这是本项目的核心工程约束之一。'))

# ---------- 五 ----------
A(('h0', '五、功能设计'))
A(('p', '按需求，四种玩法全部纳入。它们不是四个独立功能，而是同一个数据内核的四种视图。'))

A(('h1', '（一）语义图谱漫游'))
A(('p', '形态：以义类节点为地图的漫游式浏览。'))
A(('p', '交互：从首页选择一个义类，如"水"，进入后画面左右分栏。左侧展开汉语义符族，即氵部字，如江、河、湖、海、深、浅；右侧展开英语词根族，即 aqua-、hydr-、mar- 及其派生词，如 aquarium、hydrogen、marine。中间是共享的义类节点。点击任一词条进入详情，展示拆解、典故与例句。'))
A(('p', '价值：这是产品的入口，也是"啊哈体验"的发生地。用户第一次看到"三点水"和"aqua-"并列在同一个节点上时，会获得一次认知冲击。这一屏决定了产品能否被理解。'))

A(('h1', '（二）词根词缀拆解器'))
A(('p', '形态：输入任意外文单词，输出拆解结果。'))
A(('p', '交互：输入 manufacture，输出 manu（手）加 fact（做）加 ure（名词后缀），意为用手做出的东西，即制造。同时给出汉语同义类汉字对照，即扌部的打、做、造，以及相关词族，即 manual、factory、perfect 等。'))
A(('p', '技术：本地词根库做精确匹配，未命中时回退到规则分割，即前后缀剥离加词根匹配。这是工具型功能，用完即走，但它是引流的钩子——用户拆任意一个词都能得到价值。'))

A(('h1', '（三）义类词卡与间隔复习'))
A(('p', '形态：基于义类的卡片复习，而非孤立的单词卡。'))
A(('p', '核心差异：卡片背面不只有中文释义，而是词根拆解、义类归属、义符对照与典故四者合一。记忆挂钩是多重的、带结构的。'))
A(('p', '算法：间隔重复，可采用 SM-2 或其改良版本。但复习单元不完全是词，可以是义类——学完"水"义类下的一组词后，进行一次义类级的整体复习。'))
A(('p', '价值：这是留存和效果的主体。前两个功能负责让用户理解产品，这个功能负责让用户真的记住。'))

A(('h1', '（四）典故故事线'))
A(('p', '形态：按历史脉络组织的叙事内容，可阅读、可分享。示例主题如下。'))
A(('li', '1066：一场战争如何改变英语的餐桌，讲诺曼征服与 cow/beef。'))
A(('li', '两个"水"：为什么拉丁和希腊都给了英语一个"水"。'))
A(('li', '隶变：汉字的断骨手术，从古文字到今文字的分水岭。'))
A(('li', '维京人的礼物：为什么英语的"他们"是 they。'))
A(('p', '价值：内容是长期粘性来源，也是社交传播的载体。故事线可做成图文卡片，便于分享。'))

A(('h1', '（五）四者的数据流闭环'))
A(('p', '四个功能共用同一份词条库，数据流为：图谱漫游用于发现，拆解器用于查询，词卡用于记忆，故事线用于深化与留存，再回流到图谱进行再次发现。'))
A(('p', '拆解器的查询记录可以反哺词卡，查过的词自动进入复习队列；词卡的学习数据可以反哺图谱，已掌握的义类标记为完成。这个闭环是产品设计的核心，四个功能必须共享一个数据层，不能各自为政。'))

# ---------- 六 ----------
A(('h0', '六、技术方案：三端跨平台架构'))
A(('p', '需求明确为手机应用、电脑端、网页端三端齐备。这三端若分头开发，等于三份工作量、三套 bug、三次维护；因此本方案的核心原则是单套代码、三端分发。'))

A(('h1', '（一）跨端架构选型'))
A(('p', '三条可选路线对比如下。'))
A(('caption', '表8　跨端架构方案对比'))
A(('table', [
    ['方案', '技术构成', '覆盖端', '优势', '劣势'],
    ['甲：Flutter 单代码库', 'Dart 与 Flutter 框架', '手机、电脑、网页全覆盖', '移动端体验最成熟，一套代码编译六端', '网页端产物体积大、首屏偏慢；文本排版控制不如浏览器原生'],
    ['乙：Web 前端加原生壳', 'Vue 或 React 前端，Tauri 2 打包壳', '网页直接部署，电脑与手机由壳打包', '内容型应用排版最佳，迭代最快，网页端零成本', 'Tauri 移动端生态较新，安卓需自测'],
    ['丙：三端分别开发', '原生安卓、原生苹果、独立网页', '三端', '每端体验最优', '工作量约为前两者三倍，个人开发者不可行'],
]))
A(('p', '推荐方案乙。理由是本产品是内容密集型应用，绝大部分界面是文字、表格与图谱，这正是浏览器排版最擅长而 Flutter 最吃力的领域。且三端需求中网页端天然免费获得，桌面端与移动端用同一套前端加壳即可，迭代一次三端同步。'))

A(('h1', '（二）推荐架构分层'))
A(('caption', '表9　推荐技术栈分层'))
A(('table', [
    ['层', '方案', '说明'],
    ['界面层', 'Vue 3 或 React 加 TypeScript', '一套前端代码，三端复用'],
    ['网页端', '静态部署至对象存储加内容分发网络', '零额外开发，随前端同步发布'],
    ['电脑端', 'Tauri 2 打包为 Windows 与 macOS 安装包', '产物体积远小于 Electron，占用内存低'],
    ['手机端', 'Tauri 2 Mobile 或 Capacitor 打包', '复用同一前端，安卓与苹果各出一个包'],
    ['数据层', '本地结构化词库加远程接口', '常用词条本地缓存保证离线与响应速度'],
    ['后端', '轻量服务加对象存储', '负责词库分发、用户进度同步与内容更新'],
    ['内容生产', '本地 Python 脚本流水线加人工校验', '词源准确性要求高，必须人工把关'],
]))
A(('p', '需要说明一条工程边界：数据层必须与界面层和壳层完全解耦。词库以中立的 JSON 结构存储，不依赖任何前端框架；这样未来若更换壳方案，数据资产不受影响。'))

A(('h1', '（三）首屏与加载策略'))
A(('p', '本产品词条库的最终规模在数千词量级，每条含拆解、典故与例句，数据量可达数兆。因此加载策略必须从一开始就按分片设计。'))
A(('li', '首屏只加载框架与首批高频义类，保证秒开。'))
A(('li', '词库按义类切分为独立分片，进入某义类时按需加载，加载后本地缓存。'))
A(('li', '典故长文与图片走独立资源，延迟加载。'))
A(('p', '关键设计约束：词条数据不得做成一个大 JSON 文件。按义类分片天然契合产品结构，应在数据建模阶段就确定。'))

A(('h1', '（四）人工智能辅助内容生产与人工校验'))
A(('p', '内容生产是本项目最大的工作量，详见第七章。人工智能的定位是起草者而非终审者，流程为：生成词条草案，与可信来源交叉核对，人工终审并标注置信度，入库并标记为已校验。'))
A(('p', '必须坚持一条铁律：词源是硬知识，出处存疑的内容绝不入库。任何一条典故都必须能指回一个可信来源。宁可少，不可错。'))

# ---------- 七 ----------
A(('h0', '七、内容生产方案'))
A(('p', '内容是本项目真正的工作量所在，也是护城河所在。'))

A(('h1', '（一）词源实测：先量清楚要生产多少'))
A(('p', '覆盖率不能估算，必须逐词实测。本轮改用机器可读词源库 kaikki.org（维基词典的结构化导出），对高考词表 3841 词逐词判定其来源，结果如下。'))
A(('caption', '表10　高考 3500 词的词源定性实测'))
A(('table', [
    ['来源类别', '词数', '占比'],
    ['古典来源（拉丁、古希腊）', '1810', '47.1%'],
    ['罗曼斯来源（古法语等，多可追至拉丁）', '217', '5.6%'],
    ['可挂古典词根合计', '2027', '52.8%'],
    ['日耳曼底层', '1483', '38.6%'],
    ['其它来源', '4', '0.1%'],
    ['抓不到词源', '327', '8.5%'],
]))
A(('p', '抽出 51 词人工核对，与维基词典一致 51 词，判定口径可信。'))
A(('p', '按词长分层，可见越长的词越可拆，但这一趋势有上限：'))
A(('caption', '表11　可挂古典词根比例的词长分布'))
A(('table', [
    ['词长', '可挂古典词根'],
    ['1 至 3 字母', '27.8%'],
    ['4 至 6 字母', '47.5%'],
    ['7 至 9 字母', '62.9%'],
    ['10 字母以上', '60.0%'],
]))
A(('p', '这说明两点。其一，短词多为日耳曼基础词（go、eat、house），本就不难，长词才是学习者的真实痛点，而长词有六成上下可拆解，产品价值集中在长难词上。其二，比例在 7 至 9 字母达峰后趋于平缓，并非越长越可拆，因此对外不宜宣称"全覆盖"，宜表述为"专治记不住的长难词"。'))
A(('p', '更关键的一项发现来自义类分桶。将这 2027 个可挂古典词根的词按现有 55 个义类做词根匹配，仅 335 个能够归入，其余 1692 个匹配不上。人工抽检 60 个归入词，精度约为 70%，误差主要来自三字母词根的歧义：par 既是"相等"又出现在 prepare 与 part 之中，med 既是拉丁 medicus"治疗"又出现在 medius"中间"之中，此类歧义靠规则无法消解。'))
A(('p', '这项发现直接决定本项目的生产性质：义类层必须是人工校订的词典数据，不能由算法推导生成。335 这个数字只说明方向可行、缺口巨大，同时也说明义类体系必须大幅扩充——现有 55 个义类远不足以承载一本高考词表。'))
A(('p', '在此基础上，本轮把词汇范围从高考 3500 词扩到雅思，用同一条管线对雅思词表 3563 词实测，结果出人意料地好：'))
A(('caption', '表12　高考词表与雅思词表的义符命中率对比'))
A(('table', [
    ['指标', '高考 3500', '雅思词表'],
    ['可挂古典词根', '2027 词（52.8%）', '2449 词（68.7%）'],
    ['日耳曼底层', '38.6%', '25.0%'],
    ['抓不到词源', '8.5%', '6.0%'],
    ['7 至 9 字母可挂根', '62.9%', '78.9%'],
    ['10 字母以上可挂根', '60.0%', '80.1%'],
]))
A(('p', '雅思词汇的可挂古典词根比例比高考词汇高出约 16 个百分点，长词区间接近八成可拆。原因是学术英语的拉丁与希腊借词密度天然高于日常英语，而高考词表里含有大量 go、eat、house 一类日耳曼基础词。这支持一个重要的产品判断：义符机制在学术与留学类词汇上，比在高考词汇上更成立。'))
A(('p', '两份词表只重叠 1696 词（高考独有 2145 词，雅思独有 1867 词），合计约 5708 词。这说明扩充词汇范围不是把同一批词换个封面，而是真正的增量，同时也意味着内容生产量相应增加。'))

A(('h1', '（二）规模估算'))
A(('caption', '表13　内容规模估算'))
A(('table', [
    ['对象', '数量级', '说明'],
    ['义类', '150 至 300 个', '覆盖核心语义类属，视词表覆盖情况上调'],
    ['汉语义符', '100 至 200 个', '常用形旁'],
    ['外文词根', '400 至 600 个', '以英语计'],
    ['外文词缀', '约 100 个', '前缀与后缀'],
    ['首批词条', '约 1000 个', '首个完整版本的目标'],
    ['典故条目', '100 至 200 条', '每条 300 至 800 字'],
]))
A(('p', '这是一个手工可完成但耗时明显的规模。核心不在数量，而在义类映射表的质量——这张表是产品的护城河。'))

A(('h1', '（三）生产流水线'))
A(('li', '第一步，建立义类本体。'))
A(('li', '第二步，为每个义类填充汉语义符，依据《说文解字》及现代汉字学资料。'))
A(('li', '第三步，为每个义类填充外文词根，依据在线词源词典与维基词典等。'))
A(('li', '第四步，从词根反向生成词条，即每个词根下拉出其派生词。'))
A(('li', '第五步，为词条生成典故与例句。'))
A(('li', '第六步，人工校验，标注置信度，入库。'))

A(('h1', '（四）可信来源清单'))
A(('caption', '表14　内容生产的可信来源'))
A(('table', [
    ['类别', '来源'],
    ['英语词源', 'kaikki.org（维基词典结构化导出）、在线词源词典、牛津英语词典（需订阅）'],
    ['拉丁与希腊语', 'Lewis 与 Short 拉丁词典、Liddell-Scott 希腊词典'],
    ['汉字字源', '《说文解字》、汉语多功能字库、汉字全息资源应用系统'],
    ['历史背景', '诺曼征服、维京时代、文艺复兴相关史料'],
]))

A(('h1', '（五）分批策略'))
A(('p', '不宜一次性做完。按义类分批，每批十个义类，做成一个可发布的版本增量。首批十二个义类已足以验证核心机制。'))

# ---------- 八 ----------
A(('h0', '八、开发路线图'))

A(('h1', '（一）阶段划分'))
A(('caption', '表15　阶段划分与交付物'))
A(('table', [
    ['阶段', '目标', '交付物', '预估工时'],
    ['阶段零：机制验证', '验证对齐机制是否真的产生认知冲击', '交互式网页演示，含3个义类、30个词，找5至10人试用并访谈', '1 至 2 周'],
    ['阶段一：最小可用版', '跑通完整链路', '网页端上线，含图谱漫游、拆解器、词卡，首批12义类约100词', '4 至 6 周'],
    ['阶段二：三端发布', '补齐桌面端与手机端', 'Tauri 桌面安装包与手机安装包，与网页端同源', '3 至 4 周'],
    ['阶段三：内容扩充', '扩到可用的词汇量', '1000 词、100 条典故，四功能齐全', '8 至 12 周'],
    ['阶段四：打磨与增长', '体验优化与社交传播', '分享卡片、学习报告、义类成就体系', '持续'],
]))

A(('h1', '（二）关键判断点'))
A(('p', '阶段零是整个项目最重要的一步。如果交互式演示的试用者没有表现出明显的"啊哈"反应，说明对齐这个机制不成立或不够强，此时应当果断调整方向，而不是继续投入做一个完整的应用。这是本规划书建议的第一件事，也是最省钱的验证方式——用一两周时间，避免几个月白做。'))
A(('p', '另一个判断点在三端发布之前。网页端跑通后，应当先观察真实使用数据，再决定是否投入打包桌面端与手机端。若网页端的使用形态已经满足需求，桌面端可以延后。'))

# ---------- 九 ----------
A(('h0', '九、风险与对策'))

A(('h1', '（一）需求验证风险'))
A(('p', '等级：最高。本项目的核心假设，即义符与词根对齐能显著提高记忆效率，目前没有实证支持，属于待验证假设。如果这个假设不成立，例如用户虽然觉得有意思但实际记忆效果并不优于传统方法，产品价值会大幅缩水。'))
A(('p', '对策：阶段零必须做真的对照测试，而不只是收集"感觉不错"的反馈。建议设置对照组的简单记忆测试，用数据而非印象做决策。'))

A(('h1', '（二）内容准确性风险'))
A(('p', '等级：高。词源和字源是硬知识，一旦出错会严重损害产品信誉。'))
A(('p', '对策：所有内容必须指回可信来源；建立待核实标记机制；发布前抽样复核；宁可少讲，不可讲错。'))

A(('h1', '（三）跨端一致性与性能风险'))
A(('p', '等级：中。三端共用一套前端，容易出现"网页上好看、手机上错位"的问题；桌面端与移动端的打包产物也需要分别调优。'))
A(('p', '对策：界面从第一天起按响应式设计，先在手机窄屏上调通再扩展到宽屏；建立三端的固定检查清单，每次发布逐项过。'))

A(('h1', '（四）上架与合规风险'))
A(('p', '等级：中。国内移动应用需完成备案；苹果应用商店需要开发者账号；部分应用商店对教育类目有资质要求。'))
A(('caption', '表16　上架与合规事项'))
A(('table', [
    ['事项', '要求', '说明'],
    ['移动应用备案', '必须', '国内上架的应用需完成备案，否则无法发布'],
    ['苹果开发者账号', '必须', '年费约99美元，用于上架苹果应用商店'],
    ['安卓应用商店', '按各商店要求', '国内主流商店各自要求不同，需逐个适配'],
    ['内容安全', '按需', '若含用户生成内容，需接入内容安全审核'],
    ['版权', '注意', '词源数据引用需注意来源版权，优先使用开放许可来源'],
]))
A(('p', '对策：网页端不受上架限制，可作为最早的发布形态，先验证产品再投入上架流程。'))

A(('h1', '（五）平台依赖风险'))
A(('p', '等级：低至中。若完全依赖单一平台生态，政策变化会影响触达。本方案因数据层与界面层解耦，迁移成本较低。'))

A(('h1', '（六）个人开发者产能风险'))
A(('p', '等级：现实且紧迫。内容生产（1000 词加 100 条典故）加上三端开发，对个人是相当大的工作量。'))
A(('p', '对策：人工智能辅助起草加分批发布；不追求一次性完备；优先保证首批内容质量；三端按网页、桌面、手机的顺序渐进交付，每一步都有可用产出。'))

# ---------- 十 ----------
A(('h0', '十、预算与资源'))
A(('caption', '表17　预算估算'))
A(('table', [
    ['项目', '估算', '说明'],
    ['网页端部署', '每年数百元以内', '域名与对象存储，早期可用免费额度'],
    ['苹果开发者账号', '每年约99美元', '上架苹果应用商店必需'],
    ['安卓上架', '0 至数百元', '部分商店收取一次性费用'],
    ['后端服务', '早期0元', '轻量服务免费额度通常够用'],
    ['人工智能内容生成', '数百元量级', '用于词条草案生成'],
    ['开发工时', '自有投入为主', '外包则成本另计'],
    ['数据来源', '主要使用开放来源', '需订阅的词典暂不使用'],
]))

# ---------- 十一 ----------
A(('h0', '十一、待决事项'))
A(('p', '以下事项需确认后才能进入实施。'))
A(('li', '产品名：义符为暂定名，需在表3所列备选中确定。'))
A(('li', '项目落地目录：本机存放位置，当前暂用 D 盘义符目录。'))
A(('li', '跨端方案：采用推荐的方案乙（Web 前端加 Tauri 壳），还是方案甲（Flutter）。'))
A(('li', '首发形态：是否按建议先发网页端，桌面端与手机端延后。'))
A(('li', '首批义类范围：是否采用附录A的十二个义类。'))
A(('li', '义类体系规模：实测显示现有 55 个义类仅能覆盖 335 个可挂古典词根的词，'
     '与一本高考词表的需要差距很大，义类表需扩充到何规模、如何分批冻结，需先定。'))
A(('li', '是否先做阶段零的网页演示：建议做，成本最低、信息量最大。'))
A(('li', '内容来源授权：是否使用需订阅的词典，暂建议不使用。'))

# ---------- 附录 ----------
A(('pagebreak', ''))
A(('h0', '附录A　义类映射表（首批12条）'))
A(('p', '本表为产品的核心数据骨架，应在内容生产的第一步完成并冻结，后续所有词条挂靠其上。'))
A(('caption', '表18　义类映射表完整版'))
A(('table', [
    ['编号', '义类', '汉语义符', '英语词根（来源）', '典型例词'],
    ['S01', '水·液体', '氵 水', 'aqua-（拉）, hydr-（希）, mar-（拉）', 'aquarium, hydrogen, marine'],
    ['S02', '手·操作', '扌 手', 'manu-（拉）, chiro-（希）', 'manual, manufacture, chiropractor'],
    ['S03', '心·情感', '忄 心', 'cord-（拉）, psych-（希）', 'courage, cordial, psychology'],
    ['S04', '走·行进', '辶 彳 足', 'ped-（拉）, gress-（拉）, vad-（拉）', 'pedal, progress, invade'],
    ['S05', '言·言语', '讠 言', 'loqu-（拉）, dict-（拉）, log-（希）', 'eloquent, predict, dialogue'],
    ['S06', '生·生命', '生', 'bio-（希）, viv-（拉）', 'biology, survive, revive'],
    ['S07', '火·光热', '火 灬', 'ign-（拉）, photo-（希）', 'ignite, photograph'],
    ['S08', '土·大地', '土', 'terr-（拉）, geo-（希）', 'territory, geology'],
    ['S09', '石·坚硬', '石', 'lith-（希）, petr-（希）', 'monolith, petrify'],
    ['S10', '大·巨大', '大', 'magn-（拉）', 'magnify, magnitude'],
    ['S11', '小·微小', '小', 'micro-（希）', 'microscope, microbe'],
    ['S12', '目·观看', '目', 'vid-/vis-（拉）, spec-（拉）', 'visible, inspect, perspective'],
]))

A(('h0', '附录B　词条数据模型字段表'))
A(('caption', '表19　词条字段定义'))
A(('table', [
    ['字段名', '类型', '必填', '说明'],
    ['拼写', '字符串', '是', '标准拼写，唯一键'],
    ['音标', '字符串', '是', '国际音标'],
    ['义类编号', '字符串', '是', '关联附录A的义类'],
    ['词根分解', '数组', '是', '每项含部件、含义、来源语言'],
    ['汉语对应', '数组', '是', '同义类汉字示例'],
    ['典故', '文本', '否', '词源故事，可关联典故库'],
    ['来源语言', '枚举', '是', '拉丁、希腊或日耳曼'],
    ['例句', '数组', '否', '双语例句'],
    ['难度等级', '枚举', '否', '对应考纲等级'],
    ['关联词', '数组', '否', '同根词列表'],
    ['置信度', '枚举', '是', '已校验、待核实'],
]))

A(('h0', '附录C　可信来源清单'))
A(('p', '所有内容必须可溯源。清单见正文表14，此处补充使用原则。'))
A(('li', '优先使用开放许可来源，避免版权风险。'))
A(('li', '任何一条典故入库前必须指回具体来源条目，不允许凭印象补写。'))
A(('li', '存在争议的词源，标注争议并在文案中回避绝对表述。'))

A(('h0', '附录D　多语种与反向应用设想'))
A(('p', '本引擎具有天然的双向性，除汉语母语者学外语之外，还有两个延伸方向。'))
A(('p', '方向一是反向应用：对外汉语教学。同一套义类映射，反过来可以帮外国人理解汉字的义符系统。汉字难学的核心障碍之一正是义符不可感知，而本产品恰好把义符显式化了。'))
A(('p', '方向二是多语种扩展：由于义类节点与语言无关，加入法语、西班牙语几乎可以直接复用拉丁词根部分；加入日语则需要额外处理同形异义问题。'))

doc_init = None

# ============================================================
# 渲染
# ============================================================
doc = Document()

# --- 页面设置 ---
for s in doc.sections:
    s.top_margin, s.bottom_margin = Cm(3.7), Cm(3.5)
    s.left_margin, s.right_margin = Cm(2.8), Cm(2.6)

# --- Normal 基线 ---
n = doc.styles['Normal']
n.font.size = Pt(SZ)
rpr = n.element.get_or_add_rPr()
rf = rpr.get_or_add_rFonts()
rf.set(qn('w:eastAsia'), FS); rf.set(qn('w:ascii'), WEST); rf.set(qn('w:hAnsi'), WEST)

def set_run(run, ea, size=SZ, bold=False):
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), ea)
    rf.set(qn('w:ascii'), WEST); rf.set(qn('w:hAnsi'), WEST)

def set_ind(p, first_chars=None, left=None, hanging=None):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind'); pPr.append(ind)
    for k in ('w:firstLine', 'w:firstLineChars', 'w:left', 'w:leftChars',
              'w:hanging', 'w:hangingChars'):
        if ind.get(qn(k)) is not None:
            del ind.attrib[qn(k)]
    if first_chars is not None:
        ind.set(qn('w:firstLineChars'), str(first_chars))
        ind.set(qn('w:firstLine'), str(int(first_chars * 3.2)))
    if left is not None:
        ind.set(qn('w:left'), str(left)); ind.set(qn('w:leftChars'), '0')
    if hanging is not None:
        ind.set(qn('w:hanging'), str(hanging)); ind.set(qn('w:hangingChars'), '0')

def set_line(p, pt=LS, before=0, after=0):
    pf = p.paragraph_format
    pf.line_spacing = Pt(pt)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)

def set_outline(p, lvl):
    pPr = p._p.get_or_add_pPr()
    old = pPr.find(qn('w:outlineLvl'))
    if old is not None:
        pPr.remove(old)
    e = OxmlElement('w:outlineLvl')
    e.set(qn('w:val'), str(lvl))
    rpr = pPr.find(qn('w:rPr'))
    if rpr is not None:
        rpr.addprevious(e)
    else:
        pPr.append(e)

def set_page_break_before(p):
    pPr = p._p.get_or_add_pPr()
    old = pPr.find(qn('w:pageBreakBefore'))
    if old is not None:
        pPr.remove(old)
    e = OxmlElement('w:pageBreakBefore')
    st = pPr.find(qn('w:pStyle'))
    if st is not None:
        st.addnext(e)
    else:
        pPr.insert(0, e)

BOLD_RE = re.compile(r'\*\*(.+?)\*\*')
def emit_runs(p, text, ea, size=SZ, bold_all=False):
    """支持 **加粗** 内联标记。"""
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            set_run(p.add_run(text[pos:m.start()]), ea, size=size, bold=bold_all)
        set_run(p.add_run(m.group(1)), ea, size=size, bold=True)
        pos = m.end()
    if pos < len(text):
        set_run(p.add_run(text[pos:]), ea, size=size, bold=bold_all)
    if not text:
        set_run(p.add_run(''), ea, size=size, bold=bold_all)

# --- 列宽：按内容分配（短列窄、长列宽，绝不均分）---
def col_widths(rows):
    ncol = len(rows[0])
    weights = []
    for ci in range(ncol):
        mx = 0
        for r in rows:
            if ci >= len(r):
                continue
            cell = r[ci]
            w = 0.0
            for ch in cell:
                w += 1.0 if ord(ch) > 0x2E80 else 0.55
            mx = max(mx, w)
        weights.append(max(mx, 2.0))
    total = sum(weights)
    MIN = 1.6   # cm
    widths = [TEXTW * (w / total) for w in weights]
    # 抬升过窄列，等比压缩其余
    for _ in range(3):
        deficit = sum(max(0.0, MIN - w) for w in widths)
        if deficit < 0.01:
            break
        widths = [w if w >= MIN else MIN for w in widths]
        over = sum(w for w in widths if w > MIN)
        target = float(TEXTW) - MIN * sum(1 for w in widths if w <= MIN)
        if over > 0:
            widths = [w if w <= MIN else w * (target / over) for w in widths]
    s = sum(widths)
    return [int(float(TEXTW) * (w / s)) for w in widths]

def add_table(rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    ws = col_widths(rows)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.width = ws[ci]
            cp = cell.paragraphs[0]
            for r in list(cp.runs):
                r._element.getparent().remove(r._element)
            set_run(cp.add_run(val), HT if ri == 0 else FS,
                    size=TBLZ, bold=(ri == 0))
            set_ind(cp, first_chars=0)
            set_line(cp, TBL_LS)
    return t

def add_toc_field(p):
    def mkfld(kind, dirty=False):
        r = p.add_run()
        fc = OxmlElement('w:fldChar')
        fc.set(qn('w:fldCharType'), kind)
        if dirty:
            fc.set(qn('w:dirty'), 'true')
        r._r.append(fc)
    mkfld('begin', dirty=True)
    r = p.add_run()
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = ' TOC \\o "1-2" \\h \\z \\u '
    r._r.append(it)
    mkfld('separate')
    p.add_run('（在 WPS 中按 Ctrl+A 后按 F9 更新域，生成目录）')
    mkfld('end')
    for r in p.runs:
        set_run(r, FS)

# --- 主渲染 ---
li_no = 0
toc_anchor = None
title_done = False

for kind, payload in C:
    if kind == 'title':
        p = doc.add_paragraph(); emit_runs(p, payload, HZ, size=22)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_ind(p, first_chars=0); set_line(p, 34, after=6)
        title_done = True
        continue
    if kind == 'subtitle':
        p = doc.add_paragraph(); emit_runs(p, payload, KT)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_ind(p, first_chars=0); set_line(p, LS, after=18)
        continue
    if kind == 'info':
        p = doc.add_paragraph(); emit_runs(p, payload, FS)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_ind(p, first_chars=0); set_line(p, LS)
        continue
    if kind == 'blank':
        p = doc.add_paragraph(); set_ind(p, first_chars=0); set_line(p, LS)
        toc_anchor = p
        continue
    if kind == 'pagebreak':
        p = doc.add_paragraph(); set_ind(p, first_chars=0); set_line(p, LS)
        set_page_break_before(p)
        continue
    if kind == 'h0':
        li_no = 0
        p = doc.add_paragraph(); emit_runs(p, payload, HT)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_ind(p, first_chars=0); set_line(p, LS, before=14, after=8)
        set_outline(p, 0)
        continue
    if kind == 'h1':
        li_no = 0
        p = doc.add_paragraph(); emit_runs(p, payload, HT)
        set_ind(p, first_chars=200); set_line(p, LS, before=8, after=4)
        set_outline(p, 1)
        continue
    if kind == 'h2':
        li_no = 0
        p = doc.add_paragraph(); emit_runs(p, payload, KT)
        set_ind(p, first_chars=200); set_line(p, LS, before=6, after=2)
        set_outline(p, 2)
        continue
    if kind == 'p':
        p = doc.add_paragraph(); emit_runs(p, payload, FS)
        set_ind(p, first_chars=200); set_line(p, LS)
        continue
    if kind == 'li':
        li_no += 1
        p = doc.add_paragraph()
        emit_runs(p, '%d. %s' % (li_no, payload), FS)
        set_ind(p, first_chars=0, left=640, hanging=640)
        set_line(p, LS)
        continue
    if kind == 'caption':
        li_no = 0
        p = doc.add_paragraph(); emit_runs(p, payload, HT, size=TBLZ)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_ind(p, first_chars=0); set_line(p, TBL_LS, before=8, after=3)
        continue
    if kind == 'table':
        add_table(payload)
        p = doc.add_paragraph(); set_ind(p, first_chars=0); set_line(p, 10)
        continue

# --- 页脚页码「— 1 —」---
def page_footer(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    def mk(txt=None, field=False):
        r = p.add_run(txt or '')
        set_run(r, '宋体', size=14)
        if field:
            a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'begin')
            b = OxmlElement('w:instrText'); b.set(qn('xml:space'), 'preserve'); b.text = ' PAGE '
            c = OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'), 'end')
            r._r.append(a); r._r.append(b); r._r.append(c)
        return r
    mk('— '); mk(field=True); mk(' —')
    set_line(p, 14)

for s in doc.sections:
    page_footer(s)

# --- 目录（放在正文前）---
if toc_anchor is not None:
    p_title = doc.add_paragraph()
    p_title.add_run('目　录')
    p_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_ind(p_title, first_chars=0); set_line(p_title, 34, after=12)
    for r in p_title.runs:
        set_run(r, HZ, size=22)
    set_page_break_before(p_title)

    p_toc = doc.add_paragraph()
    add_toc_field(p_toc)
    set_ind(p_toc, first_chars=0); set_line(p_toc, LS)

    anchor_el = toc_anchor._p
    for el in (p_toc._p, p_title._p):
        el.getparent().remove(el)
        anchor_el.addnext(el)

    # 正文首页（第一个 h0）另起一页
    for p in doc.paragraphs:
        if p.text.strip().startswith('一、项目背景'):
            set_page_break_before(p)
            break

# --- 收尾：抹掉样式名，统一字体 ---
for p in doc.paragraphs:
    if p.style.name != 'Normal':
        p.style = doc.styles['Normal']
for p in doc.paragraphs:
    for r in p.runs:
        rpr = r._element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
        if rf.get(qn('w:eastAsia')) == '微软雅黑':
            rf.set(qn('w:eastAsia'), FS)

os.makedirs(OUT_DIR, exist_ok=True)
doc.save(DST)

# --- 清掉模板带的孤立字体引用（WPS「缺失字体」告警的两个来源）---
def patch_zip(path):
    tmp = path + '.tmp'
    zin = zipfile.ZipFile(path, 'r')
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for it in zin.infolist():
        data = zin.read(it.filename)
        name = it.filename
        if name in ('word/styles.xml', 'word/stylesWithEffects.xml'):
            s = data.decode('utf-8')
            for f in ('Courier',):
                s = s.replace('w:ascii="%s"' % f, 'w:ascii="Times New Roman"')
                s = s.replace('w:hAnsi="%s"' % f, 'w:hAnsi="Times New Roman"')
                s = s.replace('w:eastAsia="%s"' % f, 'w:eastAsia="Times New Roman"')
                s = s.replace('w:cs="%s"' % f, 'w:cs="Times New Roman"')
            s = s.replace('"微软雅黑"', '"仿宋"')
            data = s.encode('utf-8')
        elif name == 'word/theme/theme1.xml':
            s = data.decode('utf-8')
            # 脚本字体表里 26 个本机没有的字体（泰文/天城文等）→ 整表清空
            s = re.sub(r'<a:font script="[^"]*" typeface="[^"]*"\s*/>', '', s)
            data = s.encode('utf-8')
        zout.writestr(it, data)
    zin.close(); zout.close()
    os.replace(tmp, path)

patch_zip(DST)

print('OK ->', DST)
print('段落', len(doc.paragraphs), '表格', len(doc.tables))
