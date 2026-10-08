# -*- coding: utf-8 -*-
r"""公文格式 docx 渲染引擎（用户 2026-10-05 定死的那套规格）。

规格来源：D:\SRT项目\_reformat_gongwen.py
  正文 仿宋三号 / 行距固定 28 磅 / 首行缩进 2 字
  标题 华文中宋二号（文档大标题）、黑体（一级）、楷体（二级）
  表格 黑体四号表头 + 仿宋四号表体
  页码 「— 1 —」居中于页脚
  目录 用 w:outlineLvl 建大纲 + TOC 域，放在正文之前
禁用：微软雅黑、Heading 样式、「•」圆点、「★①②」序号。

内容以 (type, payload) 元组列表传入，type 取值：
  'title'     文档大标题（华文中宋二号居中）
  'subtitle'  副标题（楷体居中）
  'info'      封面信息行（仿宋居中）
  'blank'     空行（同时作为目录插入锚点）
  'pagebreak' 分页
  'h0'        一级标题（黑体居中，大纲 0 级）
  'h1'        二级标题（黑体，大纲 1 级）
  'h2'        三级标题（楷体，大纲 2 级）
  'p'         正文段（首行缩进 2 字，支持 **加粗** 内联标记）
  'li'        编号条目（自动编号，悬挂缩进）
  'caption'   表题（黑体四号居中）
  'table'     表格，payload = 二维列表，第 0 行为表头

改内容只动 C 列表，不用碰本文件的排版代码。
"""
import os, re, zipfile
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FS, HT, KT, HZ = '仿宋', '黑体', '楷体', '华文中宋'
SZ = 16          # 三号
LS = 28          # 行距固定值 28 磅
TBLZ = 14        # 表格四号
TBL_LS = 22
WEST = 'Times New Roman'
TEXTW = Cm(15.6)

BOLD_RE = re.compile(r'\*\*(.+?)\*\*')


def set_run(run, ea, size=SZ, bold=False):
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), ea)
    rf.set(qn('w:ascii'), WEST); rf.set(qn('w:hAnsi'), WEST)


def set_ind(p, first_chars=None, left=None, hanging=None):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind'); pPr.append(ind)
    for k in ('w:firstLine', 'w:firstLineChars', 'w:left', 'w:leftChars',
              'w:hanging', 'w:hangingChars'):
        if ind.get(qn(k)) is not None:
            del ind.attrib[qn(k)]
    if first_chars is not None:
        ind.set(qn('w:firstLineChars'), str(first_chars))
        ind.set(qn('w:firstLine'), str(int(first_chars * 3.2)))
    if left is not None:
        ind.set(qn('w:left'), str(left)); ind.set(qn('w:leftChars'), '0')
    if hanging is not None:
        ind.set(qn('w:hanging'), str(hanging)); ind.set(qn('w:hangingChars'), '0')


def set_line(p, pt=LS, before=0, after=0):
    pf = p.paragraph_format
    pf.line_spacing = Pt(pt)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)


def set_outline(p, lvl):
    pPr = p._p.get_or_add_pPr()
    old = pPr.find(qn('w:outlineLvl'))
    if old is not None:
        pPr.remove(old)
    e = OxmlElement('w:outlineLvl')
    e.set(qn('w:val'), str(lvl))
    rpr = pPr.find(qn('w:rPr'))
    if rpr is not None:
        rpr.addprevious(e)
    else:
        pPr.append(e)


def set_page_break_before(p):
    pPr = p._p.get_or_add_pPr()
    old = pPr.find(qn('w:pageBreakBefore'))
    if old is not None:
        pPr.remove(old)
    e = OxmlElement('w:pageBreakBefore')
    st = pPr.find(qn('w:pStyle'))
    if st is not None:
        st.addnext(e)
    else:
        pPr.insert(0, e)


def emit_runs(p, text, ea, size=SZ, bold_all=False):
    """支持 **加粗** 内联标记。"""
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            set_run(p.add_run(text[pos:m.start()]), ea, size=size, bold=bold_all)
        set_run(p.add_run(m.group(1)), ea, size=size, bold=True)
        pos = m.end()
    if pos < len(text):
        set_run(p.add_run(text[pos:]), ea, size=size, bold=bold_all)
    if not text:
        set_run(p.add_run(''), ea, size=size, bold=bold_all)


# --- 列宽：按内容分配（短列窄、长列宽，绝不均分）---
# 权重取内容宽度的 **平方根**：直接按比例会让最长的列吃掉绝大部分版面，
# 把窄列压成一字一行；开方把差距压缩到合理区间，再对过窄列做抬升。
def col_widths(rows):
    ncol = len(rows[0])
    TOTAL = float(TEXTW.cm)
    weights = []
    for ci in range(ncol):
        mx = 0.0
        for r in rows:
            if ci >= len(r):
                continue
            cell = r[ci]
            w = 0.0
            for ch in cell:
                w += 1.0 if ord(ch) > 0x2E80 else 0.55
            mx = max(mx, w)
        weights.append(max(mx, 1.0) ** 0.5)
    s = sum(weights)
    widths = [TOTAL * (w / s) for w in weights]
    MIN = 1.5   # cm，任何列不得窄于此
    for _ in range(4):
        defi = sum(MIN - w for w in widths if w < MIN)
        if defi < 0.01:
            break
        flex = sum(w - MIN for w in widths if w >= MIN)
        if flex <= 0:
            break
        widths = [MIN if w < MIN else w - (w - MIN) * defi / flex for w in widths]
    s = sum(widths)
    return [Cm(TOTAL * (w / s)) for w in widths]


def add_table(doc, rows):
    ncol = max(len(r) for r in rows)
    rows = [list(r) + [''] * (ncol - len(r)) for r in rows]   # 兜底：短行补空格
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    ws = col_widths(rows)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.width = ws[ci]
            cp = cell.paragraphs[0]
            for r in list(cp.runs):
                r._element.getparent().remove(r._element)
            set_run(cp.add_run(val), HT if ri == 0 else FS,
                    size=TBLZ, bold=(ri == 0))
            set_ind(cp, first_chars=0)
            set_line(cp, TBL_LS)
    return t


def add_toc_field(p, hint='（在 WPS 中按 Ctrl+A 后按 F9 更新域，生成目录）'):
    def mkfld(kind, dirty=False):
        r = p.add_run()
        fc = OxmlElement('w:fldChar')
        fc.set(qn('w:fldCharType'), kind)
        if dirty:
            fc.set(qn('w:dirty'), 'true')
        r._r.append(fc)
    mkfld('begin', dirty=True)
    r = p.add_run()
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = ' TOC \\o "1-2" \\h \\z \\u '
    r._r.append(it)
    mkfld('separate')
    p.add_run(hint)
    mkfld('end')
    for r in p.runs:
        set_run(r, FS)


def page_footer(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def mk(txt=None, field=False):
        r = p.add_run(txt or '')
        set_run(r, '宋体', size=14)
        if field:
            a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'begin')
            b = OxmlElement('w:instrText'); b.set(qn('xml:space'), 'preserve'); b.text = ' PAGE '
            c = OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'), 'end')
            r._r.append(a); r._r.append(b); r._r.append(c)
        return r

    mk('— '); mk(field=True); mk(' —')
    set_line(p, 14)


def patch_zip(path):
    """清掉模板带的孤立字体引用（WPS「缺失字体」告警的两个来源）。"""
    tmp = path + '.tmp'
    zin = zipfile.ZipFile(path, 'r')
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for it in zin.infolist():
        data = zin.read(it.filename)
        name = it.filename
        if name in ('word/styles.xml', 'word/stylesWithEffects.xml'):
            s = data.decode('utf-8')
            for f in ('Courier',):
                s = s.replace('w:ascii="%s"' % f, 'w:ascii="Times New Roman"')
                s = s.replace('w:hAnsi="%s"' % f, 'w:hAnsi="Times New Roman"')
                s = s.replace('w:eastAsia="%s"' % f, 'w:eastAsia="Times New Roman"')
                s = s.replace('w:cs="%s"' % f, 'w:cs="Times New Roman"')
            s = s.replace('"微软雅黑"', '"仿宋"')
            data = s.encode('utf-8')
        elif name == 'word/theme/theme1.xml':
            s = data.decode('utf-8')
            # 脚本字体表里 26 个本机没有的字体（泰文/天城文等）→ 整表清空
            s = re.sub(r'<a:font script="[^"]*" typeface="[^"]*"\s*/>', '', s)
            data = s.encode('utf-8')
        zout.writestr(it, data)
    zin.close(); zout.close()
    os.replace(tmp, path)


def build(C, dst, toc=True, toc_hint=None):
    """按 C 列表渲染成公文格式 docx，写到 dst。返回 doc。"""
    doc = Document()

    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(3.7), Cm(3.5)
        s.left_margin, s.right_margin = Cm(2.8), Cm(2.6)

    n = doc.styles['Normal']
    n.font.size = Pt(SZ)
    rpr = n.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), FS); rf.set(qn('w:ascii'), WEST); rf.set(qn('w:hAnsi'), WEST)

    li_no = 0
    toc_anchor = None
    first_h0 = None

    for kind, payload in C:
        if kind == 'title':
            p = doc.add_paragraph(); emit_runs(p, payload, HZ, size=22)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_ind(p, first_chars=0); set_line(p, 34, after=6)
            continue
        if kind == 'subtitle':
            p = doc.add_paragraph(); emit_runs(p, payload, KT)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_ind(p, first_chars=0); set_line(p, LS, after=18)
            continue
        if kind == 'info':
            p = doc.add_paragraph(); emit_runs(p, payload, FS)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_ind(p, first_chars=0); set_line(p, LS)
            continue
        if kind == 'blank':
            p = doc.add_paragraph(); set_ind(p, first_chars=0); set_line(p, LS)
            if toc_anchor is None:      # 目录锚点 = 封面后的第一个空行
                toc_anchor = p
            continue
        if kind == 'pagebreak':
            p = doc.add_paragraph(); set_ind(p, first_chars=0); set_line(p, LS)
            set_page_break_before(p)
            continue
        if kind == 'h0':
            li_no = 0
            p = doc.add_paragraph(); emit_runs(p, payload, HT)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_ind(p, first_chars=0); set_line(p, LS, before=14, after=8)
            set_outline(p, 0)
            if first_h0 is None:
                first_h0 = p
            continue
        if kind == 'h1':
            li_no = 0
            p = doc.add_paragraph(); emit_runs(p, payload, HT)
            set_ind(p, first_chars=200); set_line(p, LS, before=8, after=4)
            set_outline(p, 1)
            continue
        if kind == 'h2':
            li_no = 0
            p = doc.add_paragraph(); emit_runs(p, payload, KT)
            set_ind(p, first_chars=200); set_line(p, LS, before=6, after=2)
            set_outline(p, 2)
            continue
        if kind == 'p':
            p = doc.add_paragraph(); emit_runs(p, payload, FS)
            set_ind(p, first_chars=200); set_line(p, LS)
            continue
        if kind == 'li':
            li_no += 1
            p = doc.add_paragraph()
            emit_runs(p, '%d. %s' % (li_no, payload), FS)
            set_ind(p, first_chars=0, left=640, hanging=640)
            set_line(p, LS)
            continue
        if kind == 'caption':
            li_no = 0
            p = doc.add_paragraph(); emit_runs(p, payload, HT, size=TBLZ)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_ind(p, first_chars=0); set_line(p, TBL_LS, before=8, after=3)
            continue
        if kind == 'table':
            add_table(doc, payload)
            p = doc.add_paragraph(); set_ind(p, first_chars=0); set_line(p, 10)
            continue

    for s in doc.sections:
        page_footer(s)

    # --- 目录（放在正文前）---
    if toc and toc_anchor is not None:
        p_title = doc.add_paragraph()
        p_title.add_run('目　录')
        p_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_ind(p_title, first_chars=0); set_line(p_title, 34, after=12)
        for r in p_title.runs:
            set_run(r, HZ, size=22)
        set_page_break_before(p_title)

        p_toc = doc.add_paragraph()
        add_toc_field(p_toc, toc_hint) if toc_hint else add_toc_field(p_toc)
        set_ind(p_toc, first_chars=0); set_line(p_toc, LS)

        anchor_el = toc_anchor._p
        for el in (p_toc._p, p_title._p):
            el.getparent().remove(el)
            anchor_el.addnext(el)

    # 正文首页（第一个 h0）另起一页
    if first_h0 is not None:
        set_page_break_before(first_h0)

    # --- 收尾：抹掉样式名，统一字体 ---
    for p in doc.paragraphs:
        if p.style.name != 'Normal':
            p.style = doc.styles['Normal']
    for p in doc.paragraphs:
        for r in p.runs:
            rpr = r._element.get_or_add_rPr()
            rf = rpr.find(qn('w:rFonts'))
            if rf is None:
                rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
            if rf.get(qn('w:eastAsia')) == '微软雅黑':
                rf.set(qn('w:eastAsia'), FS)

    os.makedirs(os.path.dirname(os.path.abspath(dst)), exist_ok=True)
    doc.save(dst)
    patch_zip(dst)
    return doc